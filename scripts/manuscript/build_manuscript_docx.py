"""Build a single .docx with manuscript body + supplementary materials.

Plain-formatting docx; figures and tables embedded inline; editable in Word.
Adapted to the opx-only manuscript scope after the cleanup audit (no cpx
or twopx model claims; LEPR pairing-matrix figures dropped).
"""
from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SECTIONS_DIR = PROJECT_ROOT / 'manuscripts' / 'opx_2026' / 'text' / 'draft' / 'sections'
FIGURES_CORE = PROJECT_ROOT / 'figures' / 'core'
FIGURES_SI = PROJECT_ROOT / 'figures' / 'SI'
TABLES_DIR = PROJECT_ROOT / 'tables'
RESULTS_DIR = PROJECT_ROOT / 'results'
PREREG_DIR = PROJECT_ROOT / 'docs' / 'preregistration'
DATA_HASHES = PROJECT_ROOT / 'data' / 'hashes.json'
OUTPUT = (PROJECT_ROOT / 'manuscripts' / 'opx_2026' / 'arxiv_submission'
          / 'manuscript.docx')


# ---------------------------------------------------------------------------
# Document setup
# ---------------------------------------------------------------------------

def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0, 0, 0)
    pf = style.paragraph_format
    pf.line_spacing = 1.15
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)

    for level, size in [(1, 14), (2, 12), (3, 11)]:
        s = doc.styles[f'Heading {level}']
        s.font.name = 'Calibri'
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)


def add_page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)


def add_page_numbers(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.text = 'PAGE'
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin); run._r.append(instr); run._r.append(fld_end)

    p.add_run(' of ')

    run = p.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.text = 'NUMPAGES'
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin); run._r.append(instr); run._r.append(fld_end)


# ---------------------------------------------------------------------------
# Markdown rendering
# ---------------------------------------------------------------------------

INLINE_RE = re.compile(r'(\*\*[^\*]+\*\*|\*[^\*]+\*|`[^`]+`)')


def add_inline_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2]); run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = p.add_run(part[1:-1]); run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = p.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
        else:
            p.add_run(part)


def render_md_table(doc: Document, table_lines: list[str]) -> None:
    if len(table_lines) < 2:
        return
    rows: list[list[str]] = []
    for line in table_lines:
        if re.match(r'^\s*\|[\s\-:|]+\|\s*$', line):
            continue
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        rows.append(cells)
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    rows = [r + [''] * (n_cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(cell_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            if i == 0:
                run.bold = True


def add_markdown_section(doc: Document, md_path: Path) -> None:
    text = md_path.read_text(encoding='utf-8')
    lines = text.split('\n')

    in_code = False
    in_table = False
    table_rows: list[str] = []
    para_buffer: list[str] = []

    def flush_paragraph() -> None:
        if para_buffer:
            joined = ' '.join(para_buffer).strip()
            if joined:
                add_inline_paragraph(doc, joined)
            para_buffer.clear()

    def flush_table() -> None:
        if table_rows:
            render_md_table(doc, list(table_rows))
            table_rows.clear()

    for line in lines:
        if line.strip().startswith('```'):
            flush_paragraph(); flush_table()
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph(line)
            run = p.runs[0] if p.runs else p.add_run('')
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            continue

        if line.strip().startswith('|') and '|' in line.strip()[1:]:
            flush_paragraph()
            in_table = True
            table_rows.append(line)
            continue
        if in_table and not line.strip().startswith('|'):
            flush_table()
            in_table = False

        if line.startswith('# '):
            flush_paragraph()
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith('## '):
            flush_paragraph()
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith('### '):
            flush_paragraph()
            doc.add_heading(line[4:].strip(), level=3)
        elif line.strip() == '':
            flush_paragraph()
        else:
            para_buffer.append(line.strip())

    flush_paragraph(); flush_table()


# ---------------------------------------------------------------------------
# Figures and tables
# ---------------------------------------------------------------------------

def add_figure(doc: Document, image_path: Path, caption_text: str,
               width_inches: float = 6.0) -> None:
    if not image_path.exists():
        raise FileNotFoundError(f'Figure missing: {image_path}')
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(width_inches))
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap.add_run(caption_text)
    cap_run.italic = True
    cap_run.font.size = Pt(10)


def read_caption_sidecar(image_path: Path) -> str:
    candidates = [image_path.with_suffix('.txt')]
    base = image_path.stem
    m_full = re.match(r'(Core_\d+[a-z]?)', base)
    if m_full:
        candidates.append(
            image_path.parent / f'{m_full.group(1)}_txt_caption.txt')
    m_num = re.match(r'(Core_\d+)', base)
    if m_num:
        candidates.append(
            image_path.parent / f'{m_num.group(1)}_txt_caption.txt')
    for c in candidates:
        if c.exists():
            return c.read_text(encoding='utf-8').strip()
    return ''


def add_csv_table(doc: Document, csv_path: Path, max_rows: int | None = None,
                  caption: str | None = None) -> None:
    if not csv_path.exists():
        raise FileNotFoundError(f'Table CSV missing: {csv_path}')
    if caption:
        cap = doc.add_paragraph()
        cap_run = cap.add_run(caption)
        cap_run.italic = True
        cap_run.font.size = Pt(10)
    with open(csv_path, newline='', encoding='utf-8') as f:
        rows = list(csv.reader(f))
    if not rows:
        return
    truncated = (max_rows is not None and len(rows) > max_rows + 1)
    rows_to_render = rows[: max_rows + 1] if truncated else rows
    n_cols = len(rows[0])
    table = doc.add_table(rows=len(rows_to_render), cols=n_cols)
    table.style = 'Table Grid'
    for i, row in enumerate(rows_to_render):
        for j, val in enumerate(row[:n_cols]):
            cell = table.rows[i].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            if i == 0:
                run.bold = True
    if truncated:
        note = doc.add_paragraph()
        run = note.add_run(
            f'[Table truncated to first {max_rows} rows. '
            f'Full table at {csv_path.relative_to(PROJECT_ROOT)}]')
        run.italic = True; run.font.size = Pt(9)


# ---------------------------------------------------------------------------
# Manifest
# ---------------------------------------------------------------------------

MAIN_FIGURES = [
    ('Figure 1',  'Core_01_fig_dataset_map.png'),
    ('Figure 2',  'Core_02_fig_citation_split.png'),
    ('Figure 3',  'Core_03_fig_methods_flowchart.png'),
    ('Figure 4',  'Core_04_fig_nb04_cross_pipeline_heatmap.png'),
    ('Figure 5',  'Core_05_fig30_bias_correction_per_regime_rmse.png'),
    ('Figure 6',  'Core_06_fig31_bias_correction_residuals.png'),
    ('Figure 7',  'Core_07_fig34_bias_correction_scorecard_delta.png'),
    ('Figure 8',  'Core_08_fig45_opx_headline.png'),
    ('Figure 9a', 'Core_09a_fig_opx_regime_families.png'),
    ('Figure 9b', 'Core_09b_fig_opx_overall_families.png'),
    ('Figure 10', 'Core_10_fig_best_vs_putirka.png'),
    ('Figure 12', 'Core_12_fig_shap_winners.png'),
    ('Figure 15', 'Core_15_fig_feature_concordance.png'),
    ('Figure 16', 'Core_16_fig_classical_equivalence.png'),
    ('Figure 17', 'Core_17_fig_partial_dependence.png'),
    ('Figure 18', 'Core_18_fig_surrogate_trees.png'),
]

CORE_SUPPLEMENT_FIGS = [
    'Core_01b_fig_dataset_map_holdout.png',
    'Core_10b_fig_arcpl_bias_corrected_vs_putirka.png',
    'Core_10c_fig_shipped_vs_putirka_expetdb.png',
    'Core_10d_fig_shipped_vs_putirka_arcpl.png',
]

PREREG_FILES = [
    ('S3.1 Pressure regime pre-registration', 'p_regime_preregistration.md'),
    ('S3.2 nb03 test protocol',                'nb03_test_protocol.md'),
]


# ---------------------------------------------------------------------------
# Main build
# ---------------------------------------------------------------------------

def main() -> None:
    doc = Document()
    setup_document(doc)
    add_page_numbers(doc)

    # Front matter + abstract on page 1
    add_markdown_section(doc, SECTIONS_DIR / '00a_front_matter.md')
    add_markdown_section(doc, SECTIONS_DIR / '00_abstract.md')
    add_page_break(doc)

    body_files = [
        '01_introduction.md',
        '02_data.md',
        '03_methods.md',
        '04_results.md',
        '05_discussion.md',
        '06_conclusions.md',
        '07_data_and_code.md',
        '08_acknowledgments_and_credit.md',
        '09_cover_letter.md',
    ]
    for fname in body_files:
        add_markdown_section(doc, SECTIONS_DIR / fname)
        add_page_break(doc)

    doc.add_heading('Figures and Tables (Main Text)', level=1)

    for label, fname in MAIN_FIGURES:
        path = FIGURES_CORE / fname
        sidecar = read_caption_sidecar(path)
        caption = f'{label}. {sidecar}' if sidecar else label
        add_figure(doc, path, caption)

    doc.add_heading('Tables', level=2)

    table_1_data = [
        ['Pipeline', 'n', 'n_cit', 'P range (kbar)', 'T range (deg C)'],
        ['opx-liq',  '600',  '93',  '0-40.05', '850-1660.1'],
        ['opx-only', '1035', '123', '0-60',    '850-1656.6'],
    ]
    cap = doc.add_paragraph()
    cap_run = cap.add_run(
        'Table 1. Corpus inventory after preprocessing. n = sample count '
        'post-filters; n_cit = unique citations.')
    cap_run.italic = True; cap_run.font.size = Pt(10)
    table = doc.add_table(rows=len(table_1_data), cols=5)
    table.style = 'Table Grid'
    for i, row in enumerate(table_1_data):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Calibri'; run.font.size = Pt(10)
            if i == 0:
                run.bold = True

    add_csv_table(
        doc, TABLES_DIR / 'table_2_per_cell_winners.csv',
        caption=('Table 2. Per-cell winning model, feature set, and '
                 'aggregate test RMSE with 95% bootstrap confidence intervals.'))
    add_csv_table(
        doc, RESULTS_DIR / 'tabpfn_head_to_head.csv',
        caption=('Table 3. TabPFN v2 vs tuned-family head-to-head per '
                 '(track, target) cell. Verdict in last column.'))
    add_csv_table(
        doc, TABLES_DIR / 'table_4_bias_correction_summary.csv',
        caption=('Table 4. Bias-correction shipped form, pre/post RMSE, '
                 'percent reduction.'))

    add_page_break(doc)

    doc.add_heading('Supplementary Materials', level=1)
    doc.add_paragraph()

    doc.add_heading('S1. Supplementary Figures', level=2)
    si_idx = 1
    si_figs = sorted(p for p in FIGURES_SI.glob('*.png'))
    for path in si_figs:
        caption = read_caption_sidecar(path) or path.stem.replace('_', ' ')
        add_figure(doc, path, f'Figure S{si_idx}. {caption}')
        si_idx += 1
    for fname in CORE_SUPPLEMENT_FIGS:
        path = FIGURES_CORE / fname
        if not path.exists():
            continue
        caption = read_caption_sidecar(path) or path.stem.replace('_', ' ')
        add_figure(doc, path, f'Figure S{si_idx}. {caption}')
        si_idx += 1

    add_page_break(doc)

    doc.add_heading('S2. Supplementary Tables', level=2)
    s_idx = 1
    for csv_path in sorted(TABLES_DIR.glob('S8_*.csv')):
        doc.add_heading(
            f'Table S{s_idx}. {csv_path.stem.replace("_", " ")}',
            level=3)
        add_csv_table(doc, csv_path, max_rows=30,
                      caption=f'Source: {csv_path.relative_to(PROJECT_ROOT)}')
        s_idx += 1

    add_page_break(doc)

    doc.add_heading('S3. Pre-registration Documents', level=2)
    for title, fname in PREREG_FILES:
        doc.add_heading(title, level=3)
        path = PREREG_DIR / fname
        if path.exists():
            add_markdown_section(doc, path)
        else:
            doc.add_paragraph(f'[Source file not found: {fname}]')

    add_page_break(doc)

    doc.add_heading('S4. SHA256 Hashes', level=2)
    if DATA_HASHES.exists():
        hashes = json.loads(DATA_HASHES.read_text(encoding='utf-8'))
        if isinstance(hashes, dict) and hashes:
            cap = doc.add_paragraph()
            cap_run = cap.add_run(
                'Table S_hashes. SHA256 hashes of canonical input data files.')
            cap_run.italic = True; cap_run.font.size = Pt(10)
            entries = list(hashes.items())
            table = doc.add_table(rows=len(entries) + 1, cols=2)
            table.style = 'Table Grid'
            header = table.rows[0]
            for j, val in enumerate(['File', 'SHA256']):
                cell = header.cells[j]
                cell.text = ''
                p = cell.paragraphs[0]
                run = p.add_run(val)
                run.font.name = 'Calibri'; run.font.size = Pt(10)
                run.bold = True
            for i, (fname, sha) in enumerate(entries, start=1):
                row = table.rows[i]
                row.cells[0].text = ''
                p0 = row.cells[0].paragraphs[0]
                r0 = p0.add_run(str(fname))
                r0.font.name = 'Calibri'; r0.font.size = Pt(9)
                row.cells[1].text = ''
                p1 = row.cells[1].paragraphs[0]
                sha_str = sha if isinstance(sha, str) else json.dumps(sha)
                r1 = p1.add_run(sha_str)
                r1.font.name = 'Consolas'; r1.font.size = Pt(8)
        else:
            doc.add_paragraph('[data/hashes.json is empty or non-dict]')
    else:
        doc.add_paragraph('[data/hashes.json not found]')

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    print(f'Manuscript built: {OUTPUT}')


if __name__ == '__main__':
    main()
