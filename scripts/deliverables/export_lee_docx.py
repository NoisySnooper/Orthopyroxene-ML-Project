#!/usr/bin/env python3
"""Export the executed advisor review notebook to a Word document.

Reads:  deliverables/lee_package_20260420/00_ADVISOR_REVIEW_executed.ipynb
Writes: deliverables/lee_package_20260420/00_ADVISOR_REVIEW.docx

Design notes:
  - Markdown cells are rendered with light-weight formatting (headings,
    bold, italics, bullet lists). Nothing fancy; Dr. Lee should be able
    to edit the doc comfortably in Word.
  - Code cells' source is hidden. Only their OUTPUTS are rendered
    (figures as images, dataframes as native Word tables). She does
    not need to read Python to read the results.
  - Images decoded from the inline base64 PNGs in the executed
    notebook, so the docx is self-contained and portable.
"""
from __future__ import annotations

import base64
import io
import json
import os
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

PROJECT_ROOT = Path(__file__).parent.parent.parent.resolve()
os.chdir(PROJECT_ROOT)
sys.path.insert(0, str(PROJECT_ROOT))

import pandas as pd  # noqa: E402
from html.parser import HTMLParser  # noqa: E402

NB_IN  = PROJECT_ROOT / 'deliverables/lee_package_20260420/00_ADVISOR_REVIEW_executed.ipynb'
DOC_OUT = PROJECT_ROOT / 'deliverables/lee_package_20260420/00_ADVISOR_REVIEW.docx'

IMG_WIDTH_IN = 6.2  # fit inside 8.5-in letter with 1-in margins


# ---------------------------------------------------------------------------
# Very small markdown renderer (only features we actually use in the nb)
# ---------------------------------------------------------------------------
HEADING_RE  = re.compile(r'^(#{1,6})\s+(.*)$')
BULLET_RE   = re.compile(r'^(\s*)[-*]\s+(.*)$')
NUMBER_RE   = re.compile(r'^(\s*)\d+\.\s+(.*)$')
BOLD_RE     = re.compile(r'\*\*([^*]+)\*\*')
ITAL_RE     = re.compile(r'(?<!\*)\*([^*]+)\*(?!\*)')
CODE_RE     = re.compile(r'`([^`]+)`')


def _add_rich_runs(para, text: str):
    """Add a paragraph of text with **bold**, *italic*, `code` runs."""
    # Tokenize by the three inline marks; simple linear scan keyed on
    # whichever regex matches earliest.
    i = 0
    while i < len(text):
        best = None
        for rx, kind in ((BOLD_RE, 'b'), (ITAL_RE, 'i'), (CODE_RE, 'c')):
            m = rx.search(text, i)
            if m and (best is None or m.start() < best[1].start()):
                best = (kind, m)
        if best is None:
            para.add_run(text[i:])
            break
        kind, m = best
        if m.start() > i:
            para.add_run(text[i:m.start()])
        run = para.add_run(m.group(1))
        if kind == 'b':
            run.bold = True
        elif kind == 'i':
            run.italic = True
        elif kind == 'c':
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
        i = m.end()


def render_markdown(doc: Document, md: str) -> None:
    lines = md.splitlines()
    in_code = False
    for line in lines:
        if line.startswith('```'):
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.font.name = 'Consolas'
            r.font.size = Pt(9)
            continue
        if not line.strip():
            doc.add_paragraph()
            continue

        m_h = HEADING_RE.match(line)
        if m_h:
            level = len(m_h.group(1))
            text = m_h.group(2).strip()
            doc.add_heading(text, level=min(level, 4))
            continue

        m_b = BULLET_RE.match(line)
        if m_b:
            para = doc.add_paragraph(style='List Bullet')
            _add_rich_runs(para, m_b.group(2))
            continue

        m_n = NUMBER_RE.match(line)
        if m_n:
            para = doc.add_paragraph(style='List Number')
            _add_rich_runs(para, m_n.group(2))
            continue

        para = doc.add_paragraph()
        _add_rich_runs(para, line)


# ---------------------------------------------------------------------------
# DataFrame HTML -> native Word table
# ---------------------------------------------------------------------------
class _TableHTMLParser(HTMLParser):
    """Strip pandas-emitted HTML back into a list-of-list-of-strings."""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.rows: list[list[str]] = []
        self.cur_row: list[str] | None = None
        self.cur_cell: list[str] | None = None
        self._header_scope = False

    def handle_starttag(self, tag, attrs):
        if tag == 'tr':
            self.cur_row = []
        elif tag in ('td', 'th'):
            self.cur_cell = []
        elif tag == 'thead':
            self._header_scope = True

    def handle_endtag(self, tag):
        if tag == 'tr' and self.cur_row is not None:
            self.rows.append(self.cur_row)
            self.cur_row = None
        elif tag in ('td', 'th') and self.cur_cell is not None:
            text = ''.join(self.cur_cell).strip()
            # collapse whitespace
            text = re.sub(r'\s+', ' ', text)
            if self.cur_row is None:
                self.cur_row = []
            self.cur_row.append(text)
            self.cur_cell = None
        elif tag == 'thead':
            self._header_scope = False

    def handle_data(self, data):
        if self.cur_cell is not None:
            self.cur_cell.append(data)


def render_dataframe_html(doc: Document, html: str) -> None:
    parser = _TableHTMLParser()
    parser.feed(html)
    rows = [r for r in parser.rows if r]
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = 'Light Grid Accent 1'
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = row[j] if j < len(row) else ''
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True


def render_plain_text(doc: Document, text: str) -> None:
    text = text.rstrip()
    if not text:
        return
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r.font.size = Pt(9)


def render_png(doc: Document, b64: str) -> None:
    raw = base64.b64decode(b64)
    stream = io.BytesIO(raw)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(stream, width=Inches(IMG_WIDTH_IN))


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main() -> int:
    if not NB_IN.exists():
        print(f'ERROR: notebook missing: {NB_IN}')
        return 2
    nb = json.loads(NB_IN.read_text(encoding='utf-8'))

    doc = Document()

    # Override default body font to something readable in Word.
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    for cell in nb['cells']:
        ctype = cell['cell_type']
        src = ''.join(cell['source']) if isinstance(cell['source'], list) else cell['source']

        if ctype == 'markdown':
            render_markdown(doc, src)
            continue

        if ctype != 'code':
            continue

        # skip the code source itself; render only outputs
        for out in cell.get('outputs', []):
            ot = out.get('output_type')
            if ot == 'stream':
                continue  # drop noisy print() output
            data = out.get('data') or {}
            if 'image/png' in data:
                b64 = data['image/png']
                if isinstance(b64, list):
                    b64 = ''.join(b64)
                render_png(doc, b64)
                continue
            if 'text/html' in data:
                html = data['text/html']
                if isinstance(html, list):
                    html = ''.join(html)
                if '<table' in html:
                    render_dataframe_html(doc, html)
                    continue
            if 'text/plain' in data:
                txt = data['text/plain']
                if isinstance(txt, list):
                    txt = ''.join(txt)
                render_plain_text(doc, txt)

    doc.save(DOC_OUT)
    size_mb = DOC_OUT.stat().st_size / 1024 / 1024
    print(f'wrote {DOC_OUT}')
    print(f'size: {size_mb:.1f} MB')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
